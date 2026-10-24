# -*- coding: utf-8 -*-
"""Excel 导出器：基于 openpyxl（流式，支持百万级消息）。

依赖：openpyxl（``pip install openpyxl``）。缺少依赖时调用会抛
``qqexport.writers.base.ExcelUnavailable``，不会让整个程序崩溃。

用法::

    from qqexport.writers.excel_writer import write_chats
    files = write_chats(chats, r"D:\\QQ导出\\excel")

产物（全部落在输出目录下）：
  * ``qq_export.xlsx`` —— 合订本：**一个会话一个 sheet**（sheet 名 = 会话名，
    超 31 字符自动截断并去重），第一行是表头、冻结首行、列宽自适应、
    「内容」「附件」等长文本列自动换行。
  * ``<类型>_<会话名>.xlsx`` —— **每个会话一个独立工作簿**（正常只有一个 sheet，
    消息超过单 sheet 行数上限时拆成 ``<会话名>`` / ``<会话名> (2)`` …），
    与 txt/csv/markdown 的「一会话一文件」保持一致；同名会话自动加 ``_2``、``_3``。
  * ``index.xlsx`` —— 汇总工作簿：一张 ``汇总`` 表列出所有会话的元信息
    （会话名、类型、聊天 ID、消息数、起止时间、最后一条消息预览）。

实现要点：
  * **流式**：会话源可以是 :class:`qqexport.model.ChatStream`（**只能遍历一遍**），
    写器边读边写、写完一个会话立刻释放，绝不把全部消息物化到内存；
    文件名预留走 ``chats.metas()``（只有元数据），不会消费消息流；
  * **write_only 模式**：openpyxl 的 ``Workbook(write_only=True)``，逐行 append，
    行数据先落到临时文件，不在内存里构造巨大的单元格列表；
  * **分片**：Excel 单个 sheet 最多 1,048,576 行（含表头）。超限的会话会被拆成
    多个 sheet，名字是 ``<会话名>`` / ``<会话名> (2)`` / ``<会话名> (3)``……；
    合订本与每会话独立工作簿都支持分片。
"""

from __future__ import annotations

import os
import re
from typing import Iterable, Iterator, Sequence

from ..model import Chat
from .base import (
    APP_NAME,
    BaseWriter,
    ExcelUnavailable,
    ProgressFn,
    attachments_text,
    chat_display_name,
    chat_type_label,
    direction_label,
    excel_sheet_title,
    message_text,
    report_progress,
    sender_label,
    single_line,
    time_range_text,
)

try:  # pragma: no cover - 依赖探测
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    _HAS_OPENPYXL = True
except Exception:  # noqa: BLE001 - 任何导入异常都视为依赖不可用
    _HAS_OPENPYXL = False

XLSX_FILENAME = "qq_export.xlsx"
INDEX_FILENAME = "index.xlsx"

# 会话明细表头（与 CSV 保持一致）
HEADERS = ["会话", "类型", "消息ID", "时间", "日期", "方向", "发送者", "消息类型", "内容", "附件"]

# 汇总表表头
INDEX_HEADERS = ["会话", "类型", "聊天ID", "消息数", "开始时间", "结束时间", "最后消息"]

# 需要自动换行的列（内容列）
WRAP_COLUMNS = {9, 10}  # 1-based：内容、附件

#: Excel 单个 sheet 的最大行数（含表头行）
MAX_SHEET_ROWS = 1_048_576

#: 每个 sheet 预留一行表头
HEADER_ROWS = 1

#: 每个 sheet 最多能放的数据行数（1,048,576 - 表头）
MAX_DATA_ROWS_PER_SHEET = MAX_SHEET_ROWS - HEADER_ROWS

#: 列宽估算的采样行数。
#: write_only 模式下列宽（cols）写在 sheet XML 的**顶部**，必须在写第一行之前确定；
#: 而「先算完所有行的宽度再写」需要把整份会话的行物化到内存 —— 百万级会话不可接受。
#: 因此只采样前若干行来估算列宽。
WIDTH_SAMPLE_ROWS = 200

#: sheet 名最大长度（Excel 限制）
MAX_SHEET_TITLE_LEN = 31

HEADER_FILL = PatternFill("solid", fgColor="DCE6F1") if _HAS_OPENPYXL else None
HEADER_FONT = Font(bold=True) if _HAS_OPENPYXL else None
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center") if _HAS_OPENPYXL else None
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True) if _HAS_OPENPYXL else None


def _require_openpyxl() -> None:
    """检查 openpyxl 是否可用，不可用抛 ExcelUnavailable。"""
    if not _HAS_OPENPYXL:
        raise ExcelUnavailable("未安装 openpyxl，无法导出 Excel（pip install openpyxl）")


#: XML 1.0 不允许的控制字符（Excel 单元格写这些会直接抛
#: ``IllegalCharacterError: cannot be used in worksheets``）
_ILLEGAL_XML_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

#: Excel 单元格上限 32767 个字符；超长会触发「发现不可读取的内容」修复提示
_MAX_CELL_CHARS = 32767
_TRUNC_SUFFIX = "…（已截断）"


def _clean(value):
    """把单元格值里的非法控制字符去掉，并把超长文本截断到 Excel 上限。"""
    if isinstance(value, str):
        value = _ILLEGAL_XML_RE.sub("", value)
        if len(value) > _MAX_CELL_CHARS:
            value = value[: _MAX_CELL_CHARS - len(_TRUNC_SUFFIX)] + _TRUNC_SUFFIX
        return value
    return value


def _display_width(text) -> int:
    """估算文本的显示宽度（中文按 2 个字符宽计算）。"""
    s = "" if text is None else str(text)
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)


def _safe_len(chats) -> int | None:
    """会话数：能取到就取，取不到（比如一次性生成器）返回 None。"""
    try:
        return len(chats)
    except TypeError:
        return None


# --------------------------------------------------------------------------
# sheet 名分配
# --------------------------------------------------------------------------
class _TitleAllocator:
    """Excel sheet 名分配器（一个工作簿一个实例）。

    规则：
      * 名字合法（去掉 ``[]:*?/\\``）、不超过 31 字符 —— 由
        :func:`excel_sheet_title` 保证；
      * 同一个工作簿里不能重名（Excel 比较 sheet 名时不区分大小写）；
      * 同名会话的主名依次追加 ``_2``、``_3``（与旧版行为一致）；
      * 超长会话（消息数超过单 sheet 行数上限）拆成多个 sheet，后缀
        ``(2)``、``(3)``……，必要时继续往后找空位。
    """

    MAX_LEN = MAX_SHEET_TITLE_LEN

    def __init__(self) -> None:
        self._used: set[str] = set()

    def base(self, name: str) -> str:
        """分配一个唯一的主 sheet 名（同名加 ``_2``/``_3``）。"""
        base = excel_sheet_title(name, max_len=self.MAX_LEN)
        title = base
        counter = 2
        while title.lower() in self._used:
            suffix = "_%d" % counter
            title = base[: self.MAX_LEN - len(suffix)] + suffix
            counter += 1
        self._used.add(title.lower())
        return title

    def reserve(self, title: str) -> str:
        """占用一个确定的名字（原样使用，不追加后缀），返回它。"""
        title = excel_sheet_title(title, max_len=self.MAX_LEN)
        self._used.add(title.lower())
        return title

    def part(self, base: str, part: int) -> str:
        """分配分片 sheet 名：``base (2)``、``base (3)``……"""
        number = max(2, int(part))
        while True:
            suffix = " (%d)" % number
            stem = base[: max(1, self.MAX_LEN - len(suffix))]
            title = stem + suffix
            if title.lower() not in self._used:
                self._used.add(title.lower())
                return title
            number += 1


# --------------------------------------------------------------------------
# sheet 写入
# --------------------------------------------------------------------------
def _init_sheet(ws, widths: Sequence[int]) -> None:
    """新建 sheet 的初始化：列宽、冻结首行、表头。

    注意：write_only 模式下「列宽(cols)」与「视图(冻结首行)」都写在 sheet XML 的
    **顶部**，必须在写第一行 *之前* 设置，否则不会被保存。
    """
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    header_cells = []
    for name in HEADERS:
        cell = WriteOnlyCell(ws, value=name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        header_cells.append(cell)
    ws.append(header_cells)


def _append_row(ws, row: Sequence) -> None:
    """追加一行数据（write_only 模式）。

    只有需要换行的两列用 :class:`WriteOnlyCell`（带样式），其余列直接写原始值
    —— 百万级消息下，少创建几千万个单元格对象差别很明显。
    """
    cells = []
    for col_idx, value in enumerate(row, start=1):
        value = _clean(value)
        if col_idx in WRAP_COLUMNS:
            cell = WriteOnlyCell(ws, value=value)
            cell.alignment = WRAP_ALIGNMENT
            cells.append(cell)
        else:
            cells.append(value)
    ws.append(cells)


def _close_workbook(wb) -> None:
    """关掉工作簿。

    write_only 模式下每个 sheet 的行数据都写在临时文件里，
    正常保存时 openpyxl 会顺手清理；这里再兜一次底，避免异常退出时留下临时文件。
    """
    for ws in list(getattr(wb, "_sheets", []) or []):
        try:
            if hasattr(ws, "closed") and not ws.closed:
                ws.close()
        except Exception:  # noqa: BLE001
            pass
    try:
        wb.close()
    except Exception:  # noqa: BLE001
        pass


class _SheetStream:
    """把「行流」写进某个工作簿的一个或多个 sheet。

    Excel 单 sheet 上限 1,048,576 行（含表头），超出就自动开新 sheet：
    第一个 sheet 用主名，后面的用 ``主名 (2)``、``主名 (3)``……
    """

    def __init__(self, wb, base_title: str, titles: _TitleAllocator, widths: Sequence[int]) -> None:
        self.wb = wb
        self.base_title = base_title
        self.titles = titles
        self.widths = widths
        self.sheets: list[str] = []
        self._ws = None
        self._rows = 0
        self._part = 1

    def _open(self) -> None:
        """开一个新 sheet（写表头）。"""
        title = self.base_title if self._part == 1 else self.titles.part(self.base_title, self._part)
        self._part += 1
        ws = self.wb.create_sheet(title)
        _init_sheet(ws, self.widths)
        self._ws = ws
        self._rows = 0
        self.sheets.append(title)

    def append(self, row: Sequence) -> None:
        """追加一行（必要时先开新 sheet）。失败直接抛异常，由调用方兜底。"""
        if self._ws is None or self._rows >= MAX_DATA_ROWS_PER_SHEET:
            self._open()
        _append_row(self._ws, row)
        self._rows += 1

    def finish(self) -> None:
        """收尾：空会话也要留一个只有表头的 sheet。"""
        if self._ws is None:
            self._open()

    def drop(self) -> None:
        """把写了一半的 sheet 从工作簿里摘掉（写失败时用，避免留下半个 sheet）。"""
        for title in self.sheets:
            try:
                self.wb.remove(self.wb[title])
            except Exception:  # noqa: BLE001
                pass
        self.sheets = []
        self._ws = None


# --------------------------------------------------------------------------
# 入口
# --------------------------------------------------------------------------
def write_chats(chats: Sequence[Chat], out_dir: str, progress: ProgressFn | None = None) -> list[str]:
    """把会话写成 Excel 文件，返回写出的文件路径列表。"""
    _require_openpyxl()
    return ExcelWriter(out_dir).write(chats, progress)


class ExcelWriter(BaseWriter):
    """Excel writer 实现（流式 + 分片）。"""

    fmt = "excel"
    ext = ".xlsx"

    def write(self, chats: Iterable[Chat], progress: ProgressFn | None = None) -> list[str]:
        """把会话写盘。

        :param chats: 会话序列。可以是普通 ``list[Chat]``，也可以是
            :class:`qqexport.model.ChatStream` / ``_CountingSource`` ——
            **只能遍历一遍**，所以这里是「边迭代边写」：
            同一个会话先写进合订本，再写进它自己的独立工作簿；
            文件名在开写前用 ``chats.metas()`` 预留好，不会消费消息流。
        :param progress: 进度回调 ``progress(done, total, msg)``
        """
        _require_openpyxl()

        # 兼容：直接传入一次性迭代器（既没有 metas()、也不能重复遍历）时，
        # 先物化成列表 —— 否则 prepare_names 会把流消费光，后面什么都写不出来。
        if not hasattr(chats, "metas") and not isinstance(chats, (list, tuple)):
            chats = list(chats)

        # 预留固定名，防止与会话文件重名；并预留所有会话的原始主名
        self.reserve_names("qq_export", "index")
        self.prepare_names(chats)  # 走 metas()，不消费消息流

        total_chats = _safe_len(chats)
        state = {"chats": 0}

        def current_total() -> int:
            """进度总数：每会话 2 个产物（合订本 sheet + 独立 xlsx）+ 主工作簿 + index。"""
            if total_chats is not None:
                return total_chats * 2 + 2
            return state["chats"] * 2 + 2

        done = 0
        index_rows: list[list] = []
        chat_files: list[str] = []

        # 1) 合订本 qq_export.xlsx：流式写，一个会话一个（或多个）sheet
        main_path = os.path.join(self.out_dir, XLSX_FILENAME)
        main_titles = _TitleAllocator()
        wb = Workbook(write_only=True)  # 流式写，内存友好
        try:
            for chat in chats:  # 只能遍历一遍
                state["chats"] += 1
                name = chat_display_name(chat)
                base = main_titles.base(name)
                stream = None
                try:
                    stream = _SheetStream(wb, base, main_titles, self._estimate_widths(chat))
                    for row in self._iter_rows(chat):
                        stream.append(row)
                    stream.finish()
                except Exception as exc:  # noqa: BLE001 - 单个会话失败不影响其它 sheet
                    self.errors.append("会话「%s」写入合订本失败：%s" % (name, exc))
                    if stream is not None:
                        stream.drop()
                done += 1
                report_progress(
                    progress, done, current_total(), "已写入合订本：%s" % base
                )

                # 2) 每个会话一个独立 xlsx（与 txt/csv/markdown 一致：一会话一文件）
                path = self.chat_path(chat, ".xlsx")
                if self.safe_write(chat, path, self._write_single):
                    chat_files.append(path)
                done += 1
                report_progress(
                    progress, done, current_total(), "已导出 %s" % os.path.basename(path)
                )

                # 3) index.xlsx 的一行（汇总表用）
                try:
                    index_rows.append(self._index_row(chat))
                except Exception as exc:  # noqa: BLE001 - 汇总失败不影响其它会话
                    self.errors.append("会话「%s」写入汇总表失败：%s" % (name, exc))

            if not wb.sheetnames:
                # 一个会话都没有时，Excel 不接受「没有工作表」的工作簿 —— 补一张空表
                _init_sheet(
                    wb.create_sheet("无会话"),
                    [max(8, _display_width(h) + 2) for h in HEADERS],
                )
            try:
                wb.save(main_path)
            except Exception as exc:  # noqa: BLE001 - 合订本失败不该拖垮整个导出
                self.errors.append("合订本 %s 保存失败：%s" % (os.path.basename(main_path), exc))
                main_path = ""
        finally:
            _close_workbook(wb)
        if main_path:
            self.register(main_path)
            done += 1
            report_progress(
                progress, done, current_total(), "已导出 %s" % os.path.basename(main_path)
            )

        # 4) index.xlsx：所有会话的元信息
        index_path = os.path.join(self.out_dir, INDEX_FILENAME)
        self._write_index(index_rows, index_path)
        self.register(index_path)
        done += 1
        report_progress(
            progress, done, current_total(), "已导出 %s" % os.path.basename(index_path)
        )

        # 文件顺序与旧版保持一致：合订本、index、每会话文件
        self.files = ([main_path] if main_path else []) + [index_path] + chat_files
        return self.files

    # -- 单会话 ------------------------------------------------------------
    def _write_single(self, chat: Chat, path: str) -> None:
        """把单个会话写成一个独立的工作簿（消息超限时拆成多个 sheet）。"""
        wb = Workbook(write_only=True)
        try:
            titles = _TitleAllocator()
            base = titles.reserve(excel_sheet_title(chat_display_name(chat)))
            stream = _SheetStream(wb, base, titles, self._estimate_widths(chat))
            try:
                for row in self._iter_rows(chat):
                    stream.append(row)
                stream.finish()
            except Exception:
                stream.drop()
                raise
            wb.save(path)
        finally:
            _close_workbook(wb)

    # -- 行 ----------------------------------------------------------------
    @staticmethod
    def _iter_rows(chat: Chat) -> Iterator[list]:
        """会话的消息行（不含表头）—— 生成器，边算边写，不物化。"""
        name = chat_display_name(chat)
        kind = chat_type_label(chat)
        for msg in chat.messages:
            stamp = msg.dt.strftime("%Y-%m-%d %H:%M:%S")
            yield [
                name,
                kind,
                msg.msg_id,
                stamp,
                stamp[:10],
                direction_label(msg),
                sender_label(msg),
                msg.type_name,
                single_line(message_text(msg)),
                single_line(attachments_text(msg)),
            ]

    @classmethod
    def _estimate_widths(cls, chat: Chat) -> list[int]:
        """估算列宽（上限 60，最小 8）。

        只采样前 :data:`WIDTH_SAMPLE_ROWS` 行 —— 见该常量的说明。
        """
        widths = [max(8, _display_width(h) + 2) for h in HEADERS]
        for idx, row in enumerate(cls._iter_rows(chat)):
            if idx >= WIDTH_SAMPLE_ROWS:
                break
            for col, value in enumerate(row):
                if col >= len(widths):
                    continue
                w = _display_width(value) + 2
                if w > widths[col]:
                    widths[col] = w
        return [min(w, 60) for w in widths]

    # -- index ------------------------------------------------------------
    @staticmethod
    def _index_row(chat: Chat) -> list:
        """index.xlsx 里某个会话的一行（元信息）。"""
        msgs = getattr(chat, "messages", None) or []
        return [
            _clean(chat_display_name(chat)),
            _clean(chat_type_label(chat)),
            getattr(chat, "chat_id", "") or "",
            len(msgs),
            msgs[0].time_text if msgs else "",
            msgs[-1].time_text if msgs else "",
            _clean(single_line(getattr(chat, "preview", "") or "")),
        ]

    def _write_index(self, rows: Sequence[Sequence], path: str) -> None:
        """写 index.xlsx（会话汇总表）。"""
        wb = Workbook()
        ws = wb.active
        ws.title = "汇总"
        ws.append(INDEX_HEADERS)
        for row in rows:
            ws.append(list(row))
        for idx, name in enumerate(INDEX_HEADERS, start=1):
            cell = ws.cell(row=1, column=idx)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = HEADER_ALIGNMENT
        widths = [max(12, _display_width(h) + 2) for h in INDEX_HEADERS]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(INDEX_HEADERS)):
            for idx, cell in enumerate(row, start=1):
                w = _display_width(cell.value) + 2
                if w > widths[idx - 1]:
                    widths[idx - 1] = w
                cell.alignment = Alignment(vertical="top", wrap_text=idx in (6, 7))
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = min(width, 60)
        ws.freeze_panes = "A2"
        wb.save(path)
