# -*- coding: utf-8 -*-
"""7 种导出格式（txt/html/json/pdf/markdown/csv/excel）的端到端测试。

不依赖 pytest：可直接 ``python tests\\test_writers.py`` 运行，
也可以用 pytest 收集（函数名以 ``test_`` 开头）。

覆盖点：
  * ExportResult.ok / chats / messages 计数正确；
  * 目标文件存在且非空；
  * TXT/HTML/JSON/Markdown/CSV 里能搜到中文消息文本「你好，世界」；
  * JSON 能被 json.load 且 chats 数量正确（并能还原成 Chat 对象）；
  * CSV 第一行含「会话」且行数 = 消息数 + 1（每个会话文件）；
  * Excel 能被 openpyxl 重新 load，首个 sheet 的某格 == 预期文本；
  * Excel 额外为每个会话输出一个独立 xlsx，能被 openpyxl 读回；
  * 同名会话不会被互相覆盖（文件名自动加 _2/_3 后缀，sheet 名 ≤31 字符仍唯一）；
  * PDF 文件头是 ``%PDF`` 且大小 > 1KB；
  * 未知格式抛 ValueError；进度回调被调用；
  * ``format_status()`` 给出可用状态 + 会产出哪几类文件（供 UI 提示）。
"""

from __future__ import annotations

import csv
import json
import os
import shutil
import sys
import tempfile
import uuid

_HERE = os.path.dirname(os.path.abspath(__file__))
_SRC = os.path.join(os.path.dirname(_HERE), "src")
for _p in (_SRC, _HERE):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from qqexport import writers  # noqa: E402
from qqexport.model import (  # noqa: E402
    Attachment,
    Chat,
    ExportResult,
    Message,
)
from qqexport.writers import base as wbase  # noqa: E402
from make_fixture_chats import make_chats, make_same_name_chats  # noqa: E402

# 沙箱不允许写 %TEMP%，临时目录统一放在工程内的 _scratch 下
SCRATCH = os.path.join(_HERE, "_scratch")

# 测试里必定出现的中文消息文本
HELLO = "你好，世界"

# fixture 的规模：单聊 7 条 + 群聊 6 条 + 空会话 0 条 = 13 条，共 3 个会话
TOTAL_CHATS = 3
TOTAL_MESSAGES = 13

# 单聊会话文件名主名（单聊_<名字>）
SINGLE_PREFIX = "单聊_"
GROUP_PREFIX = "群聊_"


# --------------------------------------------------------------------------
# 工具
# --------------------------------------------------------------------------
def _tmpdir() -> str:
    """在工程内创建一个临时目录。

    先用 tempfile.mkdtemp（标准做法）；但部分受限沙箱下 mkdtemp 建出的
    目录 ACL 不允许后续写入（PermissionError），此时退回普通 makedirs。
    """
    os.makedirs(SCRATCH, exist_ok=True)
    try:
        path = tempfile.mkdtemp(prefix="writers_", dir=SCRATCH)
    except OSError:
        path = ""
    if path:
        probe = os.path.join(path, ".probe")
        try:
            with open(probe, "w", encoding="utf-8") as fh:
                fh.write("x")
            os.remove(probe)
            return path
        except OSError:
            shutil.rmtree(path, ignore_errors=True)
    path = os.path.join(SCRATCH, "writers_" + uuid.uuid4().hex[:10])
    os.makedirs(path, exist_ok=True)
    return path


def _cleanup(path: str) -> None:
    shutil.rmtree(path, ignore_errors=True)


def _read(path: str, encoding: str = "utf-8") -> str:
    with open(path, "r", encoding=encoding, newline="") as fh:
        return fh.read()


def _chat_files(out_dir: str) -> list[str]:
    """目录下的会话文件（排除 index.*）。"""
    names = []
    for name in sorted(os.listdir(out_dir)):
        if name.startswith("index"):
            continue
        if name.startswith(SINGLE_PREFIX) or name.startswith(GROUP_PREFIX):
            names.append(name)
    return names


def _file_of(chat, ext: str) -> str:
    """某个会话对应的导出文件名（<类型>_<会话名><扩展名>）。"""
    return wbase.file_stem(chat) + ext


def _export(fmt: str, chats, out_root: str, progress=None) -> ExportResult:
    res = writers.write_export(fmt, chats, out_root, progress=progress)
    assert isinstance(res, ExportResult), res
    return res


# --------------------------------------------------------------------------
# 公共断言
# --------------------------------------------------------------------------
def test_formats_table():
    """FORMATS / EXTENSIONS 表正确，7 种格式齐全。"""
    assert len(writers.FORMATS) == 7, writers.FORMATS
    assert len(writers.EXTENSIONS) == 7, writers.EXTENSIONS
    assert writers.FORMATS["txt"] == "TXT"
    assert writers.FORMATS["markdown"] == "Markdown"
    assert writers.FORMATS["excel"] == "Excel"
    assert writers.EXTENSIONS["markdown"] == ".md"
    assert writers.EXTENSIONS["excel"] == ".xlsx"
    for key in writers.FORMATS:
        assert key in writers.EXTENSIONS, key


def test_available_formats():
    """available_formats 覆盖全部无依赖格式，并给出状态说明。"""
    avail = writers.available_formats()
    for key in ("txt", "html", "json", "markdown", "csv"):
        assert key in avail, avail
    status = writers.format_status()
    assert set(status) == set(writers.FORMATS)
    # 状态说明里带上了「会产出哪几类文件」，供 UI 在导出前提示用户
    for key, text in status.items():
        assert text.startswith("可用") or text.startswith("不可用"), (key, text)
        assert "产物：" in text, (key, text)
    outputs = writers.format_outputs()
    assert set(outputs) == set(writers.FORMATS), outputs
    for key, items in outputs.items():
        assert items and all(isinstance(x, str) and x for x in items), (key, items)
    excel_desc = writers.format_output_desc("excel")
    assert "每会话" in excel_desc and excel_desc.count("xlsx") >= 2, excel_desc
    # 依赖齐全时 7 种全部可用
    if writers.is_available("pdf") and writers.is_available("excel"):
        assert len(avail) == 7, avail


def test_unknown_format_raises():
    """未知格式抛 ValueError。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        try:
            writers.write_export("docx", chats, out)
        except ValueError as exc:
            assert "docx" in str(exc), exc
        else:
            raise AssertionError("未知格式应当抛 ValueError")
    finally:
        _cleanup(out)


def test_progress_callback():
    """进度回调被调用，且 done/total 合法。"""
    chats = make_chats()
    out = _tmpdir()
    seen = []
    try:
        res = _export("txt", chats, out, progress=lambda d, t, m: seen.append((d, t, m)))
        assert res.ok, res.error
        assert seen, "进度回调没有被调用"
        for done, total, msg in seen:
            assert isinstance(done, int) and isinstance(total, int)
            assert 0 < done <= total, (done, total)
            assert isinstance(msg, str) and msg
        assert seen[-1][0] == seen[-1][1], seen[-1]
    finally:
        _cleanup(out)


def test_sanitize_filename():
    """文件名清洗规则。"""
    assert wbase.sanitize_filename('a/b:c*?"<>|d') == "abcd"
    assert wbase.sanitize_filename("  ") == "未命名会话"
    assert wbase.sanitize_filename("名字.") == "名字"
    assert wbase.sanitize_filename("a\x00b\x07c") == "abc"
    long_name = "x" * 200
    cleaned = wbase.sanitize_filename(long_name)
    assert len(cleaned) <= 120, len(cleaned)
    assert cleaned != long_name
    # 过长名字加 hash 后缀，两个不同长名不会撞成同一个
    other = wbase.sanitize_filename("y" * 200)
    assert cleaned != other


def test_empty_chat_handled():
    """空会话在所有格式里都能安全处理。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        for fmt in ("txt", "html", "markdown", "csv"):
            res = _export(fmt, chats, out)
            assert res.ok, (fmt, res.error)
            assert res.chats == TOTAL_CHATS, (fmt, res.chats)
            assert res.messages == TOTAL_MESSAGES, (fmt, res.messages)
    finally:
        _cleanup(out)


# --------------------------------------------------------------------------
# 各格式
# --------------------------------------------------------------------------
def test_txt():
    """TXT：每会话一个文件，头部有元信息，正文含中文消息。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        res = _export("txt", chats, out)
        assert res.ok, res.error
        assert res.chats == TOTAL_CHATS and res.messages == TOTAL_MESSAGES, (res.chats, res.messages)
        assert os.path.isdir(res.path), res.path
        files = _chat_files(res.path)
        assert len(files) == 3, files
        for name in files:
            path = os.path.join(res.path, name)
            assert os.path.getsize(path) > 0, path
        text = "".join(_read(os.path.join(res.path, n)) for n in files)
        assert HELLO in text, "TXT 里没有找到中文消息"
        assert "[2026-10-20 20:00:00] 我: " + HELLO in text, text[:500]
        assert "[图片] 照片.png" in text, "附件描述缺失"
        # index 文件
        index = os.path.join(res.path, "index.txt")
        assert os.path.exists(index) and os.path.getsize(index) > 0
    finally:
        _cleanup(out)


def test_html():
    """HTML：自包含、气泡左右区分、日期分隔、附件占位框、index.html。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        res = _export("html", chats, out)
        assert res.ok, res.error
        files = _chat_files(res.path)
        assert len(files) == 3, files
        for chat in chats:
            name = _file_of(chat, ".html")
            path = os.path.join(res.path, name)
            assert os.path.exists(path), path
            assert os.path.getsize(path) > 0, path
            html = _read(path)
            assert html.startswith("<!DOCTYPE html>"), name
            assert "<style>" in html, "CSS 未内联"
            assert "src=" not in html and "<link" not in html, "HTML 引用了外部资源"
            assert "date-sep" in html, "缺少日期分隔条"
            assert "file-box" in html, "缺少附件占位框"
            assert "base64" not in html, "不应嵌入 base64 图片"
        # 中文消息在单聊文件里
        single_html = _read(os.path.join(res.path, _file_of(chats[0], ".html")))
        assert HELLO in single_html, "单聊 HTML 里没有找到中文消息"
        # 我方气泡在单聊文件里（第一行消息是「我」发的）
        single_html = _read(os.path.join(res.path, _file_of(chats[0], ".html")))
        assert "row mine" in single_html, "缺少我方气泡"
        assert '<div class="row">' in single_html, "缺少对方气泡"
        # index.html
        index = os.path.join(res.path, "index.html")
        assert os.path.exists(index) and os.path.getsize(index) > 0
        index_html = _read(index)
        for name in files:
            assert name in index_html, "index.html 未列出 %s" % name
    finally:
        _cleanup(out)


def test_json():
    """JSON：单文件 qq_export.json，结构正确且能被程序读回。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        res = _export("json", chats, out)
        assert res.ok, res.error
        path = os.path.join(res.path, "qq_export.json")
        assert os.path.exists(path) and os.path.getsize(path) > 0
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        assert data["app"] == "QQ导出助手", data.get("app")
        assert data["version"] == 1
        assert "exported_at" in data
        assert data["chat_count"] == TOTAL_CHATS, data["chat_count"]
        assert data["message_count"] == TOTAL_MESSAGES, data["message_count"]
        assert len(data["chats"]) == TOTAL_CHATS, len(data["chats"])
        assert HELLO in json.dumps(data, ensure_ascii=False), "JSON 中找不到中文消息"
        # 能被程序读回：还原成 Chat / Message 对象
        restored = []
        for c in data["chats"]:
            chat = Chat(
                chat_id=c["chat_id"],
                name=c["name"],
                peer=c["peer"],
                is_group=c["is_group"],
                source=c["source"],
            )
            for m in c["messages"]:
                chat.messages.append(
                    Message(
                        msg_id=m["msg_id"],
                        time=m["time"],
                        direction=m["direction"],
                        sender=m["sender"],
                        msg_type=m["msg_type"],
                        text=m["text"],
                        content=m["content"],
                        is_mine=m["is_mine"],
                        attachments=[
                            Attachment(
                                kind=a["kind"],
                                name=a["name"],
                                size=a["size"],
                                local_path=a["local_path"],
                                md5=a["md5"],
                                url=a.get("url", ""),
                            )
                            for a in m["attachments"]
                        ],
                    )
                )
            restored.append(chat)
        assert sum(c.count for c in restored) == TOTAL_MESSAGES
        assert restored[0].messages[0].content == HELLO
    finally:
        _cleanup(out)


def test_markdown():
    """Markdown：标题层级、日期分组、引用块、附件列表。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        res = _export("markdown", chats, out)
        assert res.ok, res.error
        files = _chat_files(res.path)
        assert len(files) == 3, files
        for name in files:
            path = os.path.join(res.path, name)
            assert os.path.getsize(path) > 0, path
        md = "".join(_read(os.path.join(res.path, n)) for n in files)
        assert HELLO in md, "Markdown 里没有找到中文消息"
        assert "# 单聊 张小明" in md, md[:200]
        assert "## 2026-10-20" in md, "缺少日期分组"
        assert "### " in md, "缺少时间小标题"
        assert "- [图片] 照片.png" in md, "附件列表缺失"
        assert "> " in md, "缺少引用块"
        index = os.path.join(res.path, "index.md")
        assert os.path.exists(index) and os.path.getsize(index) > 0
    finally:
        _cleanup(out)


def test_csv():
    """CSV：带 BOM、表头含「会话」、每个会话文件行数 = 消息数 + 1。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        res = _export("csv", chats, out)
        assert res.ok, res.error
        files = _chat_files(res.path)
        assert len(files) == 3, files
        for chat in chats:
            name = _file_of(chat, ".csv")
            path = os.path.join(res.path, name)
            assert os.path.exists(path), path
            assert os.path.getsize(path) > 0, path
            raw = open(path, "rb").read()
            assert raw[:3] == b"\xef\xbb\xbf", "CSV 缺少 UTF-8 BOM"
            with open(path, "r", encoding="utf-8-sig", newline="") as fh:
                rows = list(csv.reader(fh))
            assert "会话" in rows[0], rows[0]
            assert rows[0][:2] == ["会话", "类型"], rows[0]
            # 行数 = 消息数 + 1（表头）
            assert len(rows) == chat.count + 1, (name, len(rows), chat.count)
            # 每行的「会话」列都是会话名
            for row in rows[1:]:
                assert row[0] == chat.name, row
        # 内容列含中文消息
        text = "".join(_read(os.path.join(res.path, n), encoding="utf-8-sig") for n in files)
        assert HELLO in text, "CSV 里没有找到中文消息"
        index = os.path.join(res.path, "index.csv")
        assert os.path.exists(index) and os.path.getsize(index) > 0
    finally:
        _cleanup(out)


def test_excel():
    """Excel：可被 openpyxl 读回，首个 sheet 的某格 == 预期文本。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  (跳过：未安装 openpyxl)")
        return
    chats = make_chats()
    out = _tmpdir()
    try:
        res = _export("excel", chats, out)
        assert res.ok, res.error
        main = os.path.join(res.path, "qq_export.xlsx")
        assert os.path.exists(main) and os.path.getsize(main) > 0
        wb = load_workbook(main)
        assert len(wb.sheetnames) == TOTAL_CHATS, wb.sheetnames
        ws = wb.worksheets[0]
        # sheet 名 = 会话名
        assert ws.title == chats[0].name, ws.title
        assert [c.value for c in ws[1]][:2] == ["会话", "类型"], [c.value for c in ws[1]]
        # 某个单元格正好等于预期文本
        assert ws.cell(row=2, column=9).value == HELLO, ws.cell(row=2, column=9).value
        assert ws.freeze_panes == "A2", ws.freeze_panes
        # 每个会话一个 sheet
        assert chats[1].name in wb.sheetnames, wb.sheetnames
        assert chats[2].name in wb.sheetnames, wb.sheetnames
        # index.xlsx
        index = os.path.join(res.path, "index.xlsx")
        assert os.path.exists(index) and os.path.getsize(index) > 0
        wb2 = load_workbook(index)
        assert wb2.sheetnames, wb2.sheetnames
        ws2 = wb2.worksheets[0]
        assert ws2.max_row == TOTAL_CHATS + 1, ws2.max_row  # 表头 + 3 个会话
        assert ws2.cell(row=1, column=1).value == "会话"
    finally:
        _cleanup(out)


def test_excel_per_chat_files():
    """Excel：每个会话还有一个独立的 xlsx，能被 openpyxl 读回。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  (跳过：未安装 openpyxl)")
        return
    chats = make_chats()
    out = _tmpdir()
    try:
        res = _export("excel", chats, out)
        assert res.ok, res.error
        for chat in chats:
            name = _file_of(chat, ".xlsx")
            path = os.path.join(res.path, name)
            assert os.path.exists(path), path
            assert os.path.getsize(path) > 0, path
            wb = load_workbook(path)
            # 每个独立工作簿只有一个 sheet
            assert len(wb.sheetnames) == 1, (name, wb.sheetnames)
            assert wb.sheetnames[0] == wbase.excel_sheet_title(chat.name), (name, wb.sheetnames)
            ws = wb.worksheets[0]
            assert [c.value for c in ws[1]][:2] == ["会话", "类型"], (name, [c.value for c in ws[1]])
            assert ws.max_row == chat.count + 1, (name, ws.max_row, chat.count)
            assert ws.freeze_panes == "A2", (name, ws.freeze_panes)
        # 单聊的独立 xlsx 里能读到中文消息
        wb = load_workbook(os.path.join(res.path, _file_of(chats[0], ".xlsx")))
        ws = wb.worksheets[0]
        assert ws.cell(row=2, column=9).value == HELLO, ws.cell(row=2, column=9).value
        # 合订本与每会话文件都在
        assert os.path.exists(os.path.join(res.path, "qq_export.xlsx"))
        assert os.path.exists(os.path.join(res.path, "index.xlsx"))
    finally:
        _cleanup(out)


def test_same_name_chats_not_overwritten():
    """同名会话不互相覆盖：文件名加 _2/_3 后缀，各文件内容独立。"""
    chats = make_same_name_chats()
    out = _tmpdir()
    try:
        res = _export("txt", chats, out)
        assert res.ok, res.error
        names = _chat_files(res.path)
        # 「同名会话_2」的原始主名被预留，所以第二个「同名会话」只能拿 _3
        expected = [
            "单聊_同名会话.txt",
            "单聊_同名会话_3.txt",
            "单聊_同名会话_2.txt",
            "群聊_同名会话.txt",
        ]
        for name in expected:
            assert name in names, (name, names)
        assert len(names) == 4, names
        contents = {name: _read(os.path.join(res.path, name)) for name in names}
        assert "第一个同名会话" in contents["单聊_同名会话.txt"]
        assert "第二个同名会话" in contents["单聊_同名会话_3.txt"]
        assert "名字里自带 _2 的会话" in contents["单聊_同名会话_2.txt"]
        assert "群聊里的同名会话" in contents["群聊_同名会话.txt"]
        assert len(set(contents.values())) == 4, "同名会话的文件内容被互相覆盖了"
    finally:
        _cleanup(out)


def test_same_name_chats_not_overwritten_excel():
    """同名会话的每会话 xlsx 也不互相覆盖。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  (跳过：未安装 openpyxl)")
        return
    chats = make_same_name_chats()
    out = _tmpdir()
    try:
        res = _export("excel", chats, out)
        assert res.ok, res.error
        names = [n for n in _chat_files(res.path) if n.endswith(".xlsx")]
        assert len(names) == 4, names
        texts = {}
        for name in names:
            wb = load_workbook(os.path.join(res.path, name))
            assert len(wb.sheetnames) == 1, (name, wb.sheetnames)
            texts[name] = wb.worksheets[0].cell(row=2, column=9).value
        assert len(set(texts.values())) == 4, texts
        # 合订本里 4 个 sheet，名字互不相同
        wb = load_workbook(os.path.join(res.path, "qq_export.xlsx"))
        assert len(wb.sheetnames) == 4, wb.sheetnames
        assert len(set(wb.sheetnames)) == 4, wb.sheetnames
        for title in wb.sheetnames:
            assert len(title) <= 31, title
    finally:
        _cleanup(out)


def test_pdf():
    """PDF：文件头 %PDF，大小 > 1KB。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        if not writers.is_available("pdf"):
            print("  (跳过：未安装 reportlab)")
            return
        res = _export("pdf", chats, out)
        assert res.ok, res.error
        combined = os.path.join(res.path, "qq_export.pdf")
        assert os.path.exists(combined), combined
        assert os.path.getsize(combined) > 1024, os.path.getsize(combined)
        with open(combined, "rb") as fh:
            assert fh.read(4) == b"%PDF", "PDF 文件头不对"
        files = _chat_files(res.path)
        assert len(files) == 3, files
        for name in files:
            path = os.path.join(res.path, name)
            assert os.path.getsize(path) > 1024, path
            with open(path, "rb") as fh:
                assert fh.read(4) == b"%PDF", name
    finally:
        _cleanup(out)


def test_missing_dependency_degradation():
    """缺依赖时抛明确异常（PdfUnavailable / ExcelUnavailable），不崩溃。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        if not writers.is_available("pdf"):
            try:
                writers.write_export("pdf", chats, out)
            except wbase.PdfUnavailable:
                pass
            else:
                raise AssertionError("缺 reportlab 时应当抛 PdfUnavailable")
        if not writers.is_available("excel"):
            try:
                writers.write_export("excel", chats, out)
            except wbase.ExcelUnavailable:
                pass
            else:
                raise AssertionError("缺 openpyxl 时应当抛 ExcelUnavailable")
    finally:
        _cleanup(out)


def test_single_chat_failure_does_not_break_all():
    """单个会话写失败不中断整体导出，错误写进 ExportResult.error。"""
    chats = make_chats()
    out = _tmpdir()
    try:
        res = _export("txt", chats, out)
        assert res.ok, res.error
        # 人工制造一个失败：把某个会话对象换成会抛异常的替身
        class _Boom:
            name = "坏会话"
            peer = ""
            chat_id = "x"
            is_group = False
            messages = []
            count = 0
            last_time = 0

            def to_dict(self, with_messages=True):
                raise RuntimeError("boom")

        bad = _Boom()
        res2 = writers.write_export("json", [bad], os.path.join(out, "json_fail"))
        assert res2.ok is False, "应当标记为失败"
        assert "boom" in res2.error, res2.error
    finally:
        _cleanup(out)


# --------------------------------------------------------------------------
# 运行入口
# --------------------------------------------------------------------------
def main() -> int:
    tests = [
        (name, obj)
        for name, obj in sorted(globals().items())
        if name.startswith("test_") and callable(obj)
    ]
    failed = 0
    for name, fn in tests:
        try:
            fn()
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print("FAIL %s: %s: %s" % (name, type(exc).__name__, exc))
            import traceback

            traceback.print_exc()
        else:
            print("ok", name)
    print("-" * 50)
    print("共 %d 个用例，失败 %d 个" % (len(tests), failed))
    if failed:
        print("WRITER TESTS FAILED")
        return 1
    print("all writer tests passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
