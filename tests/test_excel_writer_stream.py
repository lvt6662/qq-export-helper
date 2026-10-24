# -*- coding: utf-8 -*-
"""Excel 导出器的「流式 + 分片」测试。

不依赖 pytest：可直接 ``python tests\\test_excel_writer_stream.py`` 运行，
也可以用 pytest 收集（函数名以 ``test_`` 开头）。

覆盖点：
  * 会话流（``ChatStream`` / 一次性源）**只被遍历一遍**，每个会话的消息只加载一次
    （旧实现遍历两遍 → 大群的消息会被加载两次）；
  * 超过 Excel 单 sheet 行数上限（1,048,576）的会话被拆成多个 sheet，
    名字是 ``<会话名>`` / ``<会话名> (2)`` / ``<会话名> (3)``；
  * 合订本 ``qq_export.xlsx`` 与每会话独立 ``<类型>_<会话名>.xlsx`` 都会分片；
  * sheet 名合法（≤31 字符、不含 ``[]:*?/\\``、不重名）；
  * ``index.xlsx`` 语义不变（会话汇总表，一会话一行）；
  * 单个会话写失败只记录错误，不中断整体导出；
  * 旧接口 ``write_chats(chats, out_dir, progress)`` 与
    ``ExcelWriter(out_dir).write(chats, progress)`` 继续可用。

分片测试用「把行数上限调小」和「把写行换成计数」两种手段，**不会真的把 110 万行
落盘**（只跑通分片分支后立刻断言 sheet 数量）。
"""

from __future__ import annotations

import contextlib
import datetime as _dt
import os
import shutil
import sys
import uuid
from unittest import mock

_HERE = os.path.dirname(os.path.abspath(__file__))
_SRC = os.path.join(os.path.dirname(_HERE), "src")
for _p in (_SRC, _HERE):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from qqexport.model import Chat, ChatStream, Message  # noqa: E402
from qqexport.writers import excel_writer  # noqa: E402
from qqexport.writers.excel_writer import ExcelWriter, write_chats  # noqa: E402

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - 依赖缺失时跳过读回断言
    load_workbook = None

# 沙箱不允许写 %TEMP%，临时目录统一放在工程内的 _scratch 下
SCRATCH = os.path.join(_HERE, "_scratch")

BASE_TIME = _dt.datetime(2026, 10, 20, 20, 0, 0)


# --------------------------------------------------------------------------
# 工具
# --------------------------------------------------------------------------
def _tmpdir() -> str:
    """在工程内创建一个临时目录。"""
    os.makedirs(SCRATCH, exist_ok=True)
    path = os.path.join(SCRATCH, "excel_stream_" + uuid.uuid4().hex[:10])
    os.makedirs(path, exist_ok=True)
    return path


def _cleanup(path: str) -> None:
    shutil.rmtree(path, ignore_errors=True)


def _msg(idx: int, text: str = "消息内容") -> Message:
    """一条普通消息。"""
    return Message(
        msg_id=idx,
        time=int(BASE_TIME.timestamp()) + idx,
        direction=2,
        sender="对方",
        msg_type=2,
        text=text,
        content=text,
    )


def _chat(name: str, count: int, chat_id: str = "1", is_group: bool = True) -> Chat:
    """一个普通会话（消息条数 = count）。"""
    chat = Chat(chat_id=chat_id, name=name, peer="", is_group=is_group)
    chat.messages = [_msg(i, "第 %d 条消息" % i) for i in range(count)]
    return chat


class _FakeMessage:
    """轻量假消息：够 Excel 写器用，又不占内存（用来模拟百万级会话）。"""

    __slots__ = ("msg_id", "dt", "direction", "msg_type", "sender", "text", "content", "attachments", "is_mine")

    def __init__(self, idx: int) -> None:
        self.msg_id = idx
        self.dt = BASE_TIME
        self.direction = 2
        self.msg_type = 2
        self.sender = "对方"
        self.text = "内容 %d" % idx
        self.content = self.text
        self.attachments = []
        self.is_mine = False

    @property
    def time_text(self) -> str:
        return self.dt.strftime("%Y-%m-%d %H:%M:%S")

    @property
    def date_text(self) -> str:
        return self.dt.strftime("%Y-%m-%d")

    @property
    def type_name(self) -> str:
        return "文本"


class _FakeMessages:
    """可重复遍历的假消息序列（不占内存）。"""

    def __init__(self, count: int) -> None:
        self._count = count

    def __len__(self) -> int:
        return self._count

    def __iter__(self):
        for i in range(self._count):
            yield _FakeMessage(i)

    def __getitem__(self, idx):
        if idx < 0:
            idx += self._count
        return _FakeMessage(idx)


class _FakeChat:
    """假会话：消息按需生成。"""

    def __init__(self, name: str, count: int, chat_id: str = "1", is_group: bool = True) -> None:
        self.chat_id = chat_id
        self.name = name
        self.peer = ""
        self.is_group = is_group
        self.messages = _FakeMessages(count)

    @property
    def count(self) -> int:
        return len(self.messages)

    @property
    def preview(self) -> str:
        return "最后一条消息"


class _OneShotSource:
    """模拟 ``ChatStream`` / ``_CountingSource``：**只允许遍历一遍**。"""

    def __init__(self, chats) -> None:
        self._chats = list(chats)
        self.iter_count = 0

    def metas(self):
        """元数据（没有消息）—— 用来预留文件名。"""
        metas = []
        for c in self._chats:
            meta = Chat(chat_id=c.chat_id, name=c.name, peer=c.peer, is_group=c.is_group)
            metas.append(meta)
        return metas

    def __len__(self) -> int:
        return len(self._chats)

    def __iter__(self):
        self.iter_count += 1
        if self.iter_count > 1:
            raise AssertionError("会话流被遍历了 %d 次（必须只遍历一次）" % self.iter_count)
        for c in self._chats:
            yield c


@contextlib.contextmanager
def _small_sheet_limit(limit: int):
    """临时把「每个 sheet 的数据行上限」调小（用来快速验证分片逻辑）。"""
    with mock.patch.object(excel_writer, "MAX_DATA_ROWS_PER_SHEET", limit):
        yield


@contextlib.contextmanager
def _count_rows_only():
    """把「写行」换成计数：验证分片逻辑，但不把上百万行真的写进 xlsx。"""
    counter = {"rows": 0, "per_sheet": {}, "titles": {}}

    def fake_append(ws, row):
        counter["rows"] += 1
        key = id(ws)  # 按 sheet 对象计数：合订本与独立 xlsx 的同名 sheet 分开算
        counter["per_sheet"][key] = counter["per_sheet"].get(key, 0) + 1
        counter["titles"][key] = ws.title

    with mock.patch.object(excel_writer, "_append_row", fake_append):
        yield counter


# --------------------------------------------------------------------------
# 基础行为
# --------------------------------------------------------------------------
def test_stream_source_is_consumed_once():
    """会话流只被遍历一遍；文件名预留走 metas()（不消费消息）。"""
    chats = [_chat("群甲", 3, "1"), _chat("群乙", 4, "2")]
    source = _OneShotSource(chats)
    out = _tmpdir()
    try:
        writer = ExcelWriter(out)
        files = writer.write(source)
        assert source.iter_count == 1, "会话流被遍历了 %d 次" % source.iter_count
        assert not writer.errors, writer.errors
        # 合订本 + index + 每会话文件
        assert len(files) == 2 + 2, files
        assert os.path.exists(os.path.join(out, "qq_export.xlsx"))
        assert os.path.exists(os.path.join(out, "index.xlsx"))
        if load_workbook is not None:
            wb = load_workbook(os.path.join(out, "qq_export.xlsx"))
            assert wb.sheetnames == ["群甲", "群乙"], wb.sheetnames
            # 消息确实写进去了（不是只写了元数据）
            ws = wb["群甲"]
            assert ws.max_row == 4, ws.max_row  # 表头 + 3 条
            assert ws.cell(row=2, column=9).value == "第 0 条消息"
    finally:
        _cleanup(out)


def test_chatstream_loader_called_once_per_chat():
    """``ChatStream`` 的 loader 每个会话只调用一次（消息只加载一次）。"""
    loaded = []

    def loader(meta: Chat) -> Chat:
        loaded.append(meta.chat_id)
        chat = Chat(chat_id=meta.chat_id, name=meta.name, peer=meta.peer, is_group=meta.is_group)
        chat.messages = [_msg(i, "消息 %d" % i) for i in range(5)]
        return chat

    metas = [Chat(chat_id="a", name="甲", is_group=True), Chat(chat_id="b", name="乙", is_group=False)]
    stream = ChatStream(metas, loader)
    out = _tmpdir()
    try:
        writer = ExcelWriter(out)
        writer.write(stream)
        assert loaded == ["a", "b"], loaded  # 每个会话只加载一次
        assert not writer.errors, writer.errors
        if load_workbook is not None:
            wb = load_workbook(os.path.join(out, "qq_export.xlsx"))
            assert wb.sheetnames == ["甲", "乙"], wb.sheetnames
    finally:
        _cleanup(out)


def test_index_semantics_unchanged():
    """index.xlsx 仍是「会话汇总表」：表头 + 每个会话一行。"""
    chats = [_chat("汇总甲", 3, "1001"), _chat("汇总乙", 0, "1002", is_group=False)]
    out = _tmpdir()
    try:
        ExcelWriter(out).write(chats)
        path = os.path.join(out, "index.xlsx")
        assert os.path.exists(path)
        if load_workbook is None:
            return
        wb = load_workbook(path)
        ws = wb.worksheets[0]
        assert ws.title == "汇总", ws.title
        assert [c.value for c in ws[1]] == excel_writer.INDEX_HEADERS
        assert ws.max_row == 3, ws.max_row  # 表头 + 2 个会话
        assert ws.cell(row=2, column=1).value == "汇总甲"
        assert ws.cell(row=2, column=3).value == "1001"
        assert ws.cell(row=2, column=4).value == 3  # 消息数
        assert ws.cell(row=3, column=4).value == 0  # 空会话
    finally:
        _cleanup(out)


def test_progress_reported():
    """进度回调被调用，且最后一次 done == total。"""
    chats = [_chat("进度甲", 2, "1"), _chat("进度乙", 2, "2")]
    seen = []
    out = _tmpdir()
    try:
        ExcelWriter(out).write(chats, progress=lambda d, t, m: seen.append((d, t, m)))
        assert seen, "进度回调没有被调用"
        for done, total, msg in seen:
            assert isinstance(done, int) and isinstance(total, int)
            assert 0 < done <= total, (done, total)
            assert isinstance(msg, str) and msg
        assert seen[-1][0] == seen[-1][1], seen[-1]
    finally:
        _cleanup(out)


def test_backward_compatible_api():
    """旧接口继续可用：write_chats(chats, out_dir, progress) / ExcelWriter(out).write(...)。"""
    chats = [_chat("兼容甲", 3, "1"), _chat("兼容乙", 4, "2")]
    out = _tmpdir()
    try:
        seen = []
        files = write_chats(chats, out, progress=lambda d, t, m: seen.append(d))
        assert seen
        assert os.path.exists(os.path.join(out, "qq_export.xlsx"))
        assert os.path.exists(os.path.join(out, "index.xlsx"))
        # 每会话一个 xlsx
        names = [n for n in os.listdir(out) if n.endswith(".xlsx") and not n.startswith("index")]
        assert len(names) == 3, names  # 合订本 + 2 个会话
        assert len(files) == 4, files  # 合订本 + index + 2 个会话
        # 再用 ExcelWriter 直接写一次
        out2 = _tmpdir()
        try:
            w = ExcelWriter(out2)
            got = w.write(chats, None)
            assert len(got) == 4, got
            assert not w.errors, w.errors
        finally:
            _cleanup(out2)
    finally:
        _cleanup(out)


def test_single_chat_failure_does_not_break_export():
    """单个会话写失败只记录错误，其它会话照常导出。"""

    class _BoomMessage:
        msg_id = 1
        direction = 2
        msg_type = 2
        sender = "坏消息"
        text = "boom"
        content = "boom"
        attachments = []
        is_mine = False

        @property
        def dt(self):
            raise RuntimeError("boom")

        @property
        def time_text(self):
            return "2026-10-20 20:00:00"

        @property
        def type_name(self):
            return "文本"

    good = _chat("好会话", 3, "1")
    bad = Chat(chat_id="2", name="坏会话", is_group=True)
    bad.messages = [_msg(1), _BoomMessage(), _msg(3)]
    out = _tmpdir()
    try:
        writer = ExcelWriter(out)
        writer.write([good, bad])
        assert writer.errors, "写失败应该有错误记录"
        assert any("boom" in e for e in writer.errors), writer.errors
        # 好会话照样写出来了
        assert os.path.exists(os.path.join(out, "qq_export.xlsx"))
        if load_workbook is not None:
            wb = load_workbook(os.path.join(out, "qq_export.xlsx"))
            assert wb.sheetnames == ["好会话"], wb.sheetnames
    finally:
        _cleanup(out)


# --------------------------------------------------------------------------
# 分片
# --------------------------------------------------------------------------
def test_sheet_limit_constants():
    """分片阈值就是 Excel 的硬上限：1,048,576 行（含表头）。"""
    assert excel_writer.MAX_SHEET_ROWS == 1048576
    assert excel_writer.MAX_DATA_ROWS_PER_SHEET == 1048575
    # 110 万条消息的会话 → 2 个 sheet
    assert -(-1100000 // excel_writer.MAX_DATA_ROWS_PER_SHEET) == 2


def test_split_into_multiple_sheets():
    """超过单 sheet 行数上限时拆成多个 sheet，名字是 ``X`` / ``X (2)`` / ``X (3)``。"""
    chat = _chat("电脑互助群", 12, "1")
    out = _tmpdir()
    try:
        with _small_sheet_limit(5):
            writer = ExcelWriter(out)
            writer.write([chat])
            assert not writer.errors, writer.errors
        if load_workbook is None:
            return
        # 合订本：3 个 sheet
        wb = load_workbook(os.path.join(out, "qq_export.xlsx"))
        assert wb.sheetnames == ["电脑互助群", "电脑互助群 (2)", "电脑互助群 (3)"], wb.sheetnames
        assert wb["电脑互助群"].max_row == 6  # 表头 + 5 行
        assert wb["电脑互助群 (2)"].max_row == 6
        assert wb["电脑互助群 (3)"].max_row == 3  # 表头 + 2 行
        # 每会话独立 xlsx：同样是 3 个 sheet
        path = os.path.join(out, "群聊_电脑互助群.xlsx")
        assert os.path.exists(path), os.listdir(out)
        wb2 = load_workbook(path)
        assert wb2.sheetnames == ["电脑互助群", "电脑互助群 (2)", "电脑互助群 (3)"], wb2.sheetnames
        assert wb2["电脑互助群 (3)"].max_row == 3
    finally:
        _cleanup(out)


def test_huge_chat_split_count():
    """110 万条消息的会话 → 合订本与独立 xlsx 各拆成 2 个 sheet（不落盘 110 万行）。"""
    big = _FakeChat("超大群", 1100000, "1")
    out = _tmpdir()
    try:
        with _count_rows_only() as counter:
            writer = ExcelWriter(out)
            writer.write([big])
            assert not writer.errors, writer.errors
        # 每个 sheet 的数据行数：满的那个 = 上限，剩下的进第二个 sheet
        assert counter["rows"] == 2 * 1100000, counter["rows"]  # 合订本 + 独立 xlsx
        by_title: dict[str, list[int]] = {}
        for key, n in counter["per_sheet"].items():
            by_title.setdefault(counter["titles"][key], []).append(n)
        # 合订本与独立 xlsx 各拆成 2 个 sheet，行数一样
        assert sorted(by_title["超大群"]) == [1048575, 1048575], by_title
        assert sorted(by_title["超大群 (2)"]) == [51425, 51425], by_title
        if load_workbook is None:
            return
        wb = load_workbook(os.path.join(out, "qq_export.xlsx"))
        assert wb.sheetnames == ["超大群", "超大群 (2)"], wb.sheetnames
        wb2 = load_workbook(os.path.join(out, "群聊_超大群.xlsx"))
        assert wb2.sheetnames == ["超大群", "超大群 (2)"], wb2.sheetnames
    finally:
        _cleanup(out)


def test_sheet_names_are_legal_and_unique():
    """sheet 名：≤31 字符、不含 ``[]:*?/\\``、不重名。"""
    long_name = "这是一个非常非常非常长的群名字" * 3
    chats = [
        _chat(long_name, 12, "1"),
        _chat("电脑互助群", 12, "2"),
        _chat("电脑互助群 (2)", 1, "3"),
        _chat("带*非法?字符/的群", 1, "4"),
    ]
    out = _tmpdir()
    try:
        with _small_sheet_limit(5):
            writer = ExcelWriter(out)
            writer.write(chats)
            assert not writer.errors, writer.errors
        if load_workbook is None:
            return
        wb = load_workbook(os.path.join(out, "qq_export.xlsx"))
        titles = wb.sheetnames
        assert len(set(titles)) == len(titles), titles  # 不重名
        for t in titles:
            assert len(t) <= 31, (t, len(t))
            for bad in "[]:*?/\\":
                assert bad not in t, (t, bad)
        # 第一个会话超长 → 主名被截断，分片名带 " (2)" 后缀
        assert titles[0].startswith("这是一个非常非常非常长的群名字"), titles[0]
        assert titles[1] == " (2)" or titles[1].endswith(" (2)"), titles[1]
    finally:
        _cleanup(out)


def test_same_name_chats_get_distinct_sheets():
    """同名会话的 sheet 名互不相同。"""
    chats = [_chat("同名群", 2, "1"), _chat("同名群", 2, "2"), _chat("同名群", 2, "3")]
    out = _tmpdir()
    try:
        ExcelWriter(out).write(chats)
        if load_workbook is None:
            return
        wb = load_workbook(os.path.join(out, "qq_export.xlsx"))
        assert len(wb.sheetnames) == 3, wb.sheetnames
        assert len(set(wb.sheetnames)) == 3, wb.sheetnames
        for t in wb.sheetnames:
            assert len(t) <= 31, t
    finally:
        _cleanup(out)


def test_write_export_with_stream_source():
    """走完整入口：``write_export("excel", 流式源, 目录)``。"""
    from qqexport import writers

    metas = [
        Chat(chat_id="a", name="流甲", is_group=True),
        Chat(chat_id="b", name="流乙", is_group=False),
    ]

    def loader(meta: Chat) -> Chat:
        chat = Chat(chat_id=meta.chat_id, name=meta.name, peer=meta.peer, is_group=meta.is_group)
        chat.messages = [_msg(i, "流消息 %d" % i) for i in range(4)]
        return chat

    stream = ChatStream(metas, loader)
    out = _tmpdir()
    try:
        res = writers.write_export("excel", stream, out)
        assert res.ok, res.error
        # 会话数 / 消息数不再被「遍历两遍」翻倍
        assert res.chats == 2, res.chats
        assert res.messages == 8, res.messages
        excel_dir = os.path.join(out, "excel")
        assert os.path.exists(os.path.join(excel_dir, "qq_export.xlsx"))
        assert os.path.exists(os.path.join(excel_dir, "index.xlsx"))
    finally:
        _cleanup(out)


def test_empty_chat_list_still_writes_valid_workbook():
    """没有会话时，合订本也要是一个「Excel 能打开」的工作簿（至少一张表）。"""
    out = _tmpdir()
    try:
        writer = ExcelWriter(out)
        writer.write([])
        assert not writer.errors, writer.errors
        path = os.path.join(out, "qq_export.xlsx")
        assert os.path.exists(path)
        if load_workbook is not None:
            wb = load_workbook(path)
            assert wb.sheetnames, "空导出也应当有一个工作表"
    finally:
        _cleanup(out)


def test_empty_chat_still_gets_sheet():
    """空会话也要有一个（只有表头的）sheet。"""
    out = _tmpdir()
    try:
        ExcelWriter(out).write([_chat("空群", 0, "1")])
        if load_workbook is None:
            return
        wb = load_workbook(os.path.join(out, "qq_export.xlsx"))
        assert wb.sheetnames == ["空群"], wb.sheetnames
        assert wb["空群"].max_row == 1, wb["空群"].max_row  # 只有表头
    finally:
        _cleanup(out)


# --------------------------------------------------------------------------
# 运行入口
# --------------------------------------------------------------------------
def main() -> int:
    tests = [
        (name, obj)
        for name, obj in sorted(globals().items())
        if name.startswith("test_") and callable(obj)
    ]
    failed = 0
    for name, fn in tests:
        try:
            fn()
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print("FAIL %s: %s: %s" % (name, type(exc).__name__, exc))
            import traceback

            traceback.print_exc()
        else:
            print("ok", name)
    print("-" * 50)
    print("共 %d 个用例，失败 %d 个" % (len(tests), failed))
    if failed:
        print("EXCEL STREAM TESTS FAILED")
        return 1
    print("all excel stream tests passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
